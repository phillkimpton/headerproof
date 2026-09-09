"""Offline acceptance tests; network and scanner failures are controlled fixtures."""
import builtins
import contextlib
import csv
import datetime
import io
import json
from pathlib import Path
import socket
import ssl
import subprocess
import sys
import tempfile
import threading
import unittest
from types import SimpleNamespace
from unittest.mock import Mock, patch

import headerproof as hp

STAMP = datetime.datetime(2026, 1, 1, tzinfo=datetime.timezone.utc)


class FixedDatetime(datetime.datetime):
    @classmethod
    def now(cls, tz=None):
        return STAMP



def scanned(target='example.com'):
    headers = {h: {'defined': True, 'warn': False, 'contents': 'value', 'notes': []}
               for h, _ in hp.MATRIX}
    headers['content-security-policy']['defined'] = False
    headers['content-security-policy']['contents'] = None
    headers['x-frame-options']['warn'] = True
    headers['server'] = {'defined': True, 'warn': False, 'contents': 'example', 'notes': []}
    return {'target': target, 'headers': headers,
            'https': {'supported': True, 'certvalid': True, 'redirect': False}}


class Acceptance(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def run_cli(self, extra=(), quiet=False, missing_xlsx=False):
        self.root.mkdir(exist_ok=True)
        stdout, stderr = io.StringIO(), io.StringIO()
        statuses = {'example.com': ('live', 'resolves', '200'),
                    'https://example.com/error': ('live', 'resolves', '403'),
                    'github.com': ('dead', 'NODNS', '000')}
        def scan(target, timeout):
            return scanned(target) if target == 'example.com' else {'target': target, 'error': 'challenge'}
        real_import = builtins.__import__
        def importer(name, *args, **kwargs):
            if missing_xlsx and name.startswith('openpyxl'):
                raise ImportError('controlled missing dependency')
            return real_import(name, *args, **kwargs)
        args = list(statuses) + ['--outdir', str(self.root), '--name', 'report'] + list(extra)
        if quiet:
            args.append('-q')
        with patch.object(hp, 'check_live', side_effect=lambda t, timeout: statuses[t]), \
             patch.object(hp, 'scan_one', side_effect=scan) as scanner, \
             patch.object(hp, 'datetime', SimpleNamespace(datetime=FixedDatetime, timezone=datetime.timezone)), \
             patch('builtins.__import__', side_effect=importer), \
             contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
            hp.main(args)
        return stdout.getvalue(), stderr.getvalue(), scanner

    def test_01_input_order_and_duplicates(self):
        path = self.root / 'targets.txt'
        path.write_text('\n# comment\nexample.com\nexample.com\nhttps://github.com/path\n')
        self.assertEqual(hp.read_targets(path), ['example.com', 'https://github.com/path'])
        with patch.object(hp, 'check_live', return_value=('dead', 'NODNS', '000')) as gate:
            hp.main(['example.com', '-l', str(path), '-q', '--outdir', str(self.root)])
        self.assertEqual([c.args[0] for c in gate.call_args_list], ['example.com', 'https://github.com/path'])

    def test_02_no_targets_usage_error(self):
        err = io.StringIO()
        with contextlib.redirect_stderr(err), self.assertRaises(SystemExit) as caught:
            hp.main([])
        self.assertNotEqual(caught.exception.code, 0)
        self.assertIn('usage:', err.getvalue())

    def test_03_resolving_http_response(self):
        for status in (200, 301, 403, 503):
            with self.subTest(status=status), patch.object(hp.socket, 'getaddrinfo') as dns, \
                 patch.object(hp.http.client, 'HTTPSConnection') as connection:
                connection.return_value.getresponse.return_value.status = status
                self.assertEqual(hp.check_live('https://example.com:8443/path?q=x'), ('live', 'resolves', str(status)))
                dns.assert_called_once_with('example.com', 8443)
                connection.return_value.request.assert_called_once_with('HEAD', '/')
                self.assertEqual(connection.call_args.kwargs['context'].verify_mode, ssl.CERT_NONE)
                connection.return_value.close.assert_called_once()

    def test_04_no_dns(self):
        with patch.object(hp.socket, 'getaddrinfo', side_effect=socket.gaierror), \
             patch.object(hp.http.client, 'HTTPSConnection') as connection:
            self.assertEqual(hp.check_live('example.com'), ('dead', 'NODNS', '000'))
            connection.assert_not_called()

    def test_05_no_https_response(self):
        for error in (TimeoutError(), ConnectionRefusedError(), ssl.SSLError(), hp.http.client.BadStatusLine('bad')):
            with self.subTest(error=type(error)), patch.object(hp.socket, 'getaddrinfo'), \
                 patch.object(hp.http.client, 'HTTPSConnection') as connection:
                connection.return_value.request.side_effect = error
                self.assertEqual(hp.check_live('example.com'), ('dead', 'resolves', '000'))
                connection.return_value.close.assert_called_once()

    def test_06_scan_and_problem_count(self):
        with patch.object(hp.subprocess, 'run', return_value=Mock(returncode=0, stdout=json.dumps(scanned()), stderr='')):
            record = hp.scan_one('example.com')
        self.assertEqual(hp.build([record])[0]['_problems'], 2)
        self.assertFalse(record['_scanerror'])

    def test_07_scanner_failure_variants(self):
        cases = [(1, '', 'denied', 'denied'), (0, '', '', 'scan failed'),
                 (0, 'not json', '', 'unparseable scanner output')]
        for code, stdout, stderr, reason in cases:
            with patch.object(hp.subprocess, 'run', return_value=Mock(returncode=code, stdout=stdout, stderr=stderr)):
                record = hp.scan_one('example.com')
            self.assertEqual(record['error'], reason)
            self.assertTrue(hp.build([record])[0]['_scanerror'])
        with patch.object(hp.subprocess, 'run', side_effect=subprocess.TimeoutExpired('secheaders', 30)):
            self.assertEqual(hp.scan_one('example.com')['error'], 'scan timeout')

    def test_08_dead_never_scanned(self):
        _, _, scanner = self.run_cli()
        self.assertEqual({c.args[0] for c in scanner.call_args_list}, {'example.com', 'https://example.com/error'})
        with patch.object(hp, 'check_live', return_value=('dead', 'resolves', '000')), \
             patch.object(hp.subprocess, 'run') as subprocess_run:
            hp.main(['github.com', '-q', '--outdir', str(self.root)])
            subprocess_run.assert_not_called()

    def test_09_path_resolution(self):
        executable = str(self.root / 'secheaders')
        Path(executable).touch()
        with patch.object(hp.shutil, 'which', return_value=executable):
            self.assertEqual(hp._secheaders_cmd(), [executable])

    def test_10_module_resolution(self):
        with patch.object(hp.shutil, 'which', return_value=None):
            self.assertEqual(hp._secheaders_cmd(), [sys.executable, '-m', 'secheaders'])

    def test_11_all_five_files(self):
        self.run_cli(['-o', 'all'])
        self.assertEqual({p.name for p in self.root.iterdir()}, {'report.' + f for f in hp.RENDERERS})

    def test_12_csv_three_states(self):
        self.run_cli(['-o', 'csv'])
        with (self.root / 'report.csv').open() as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual([r['scan_status'] for r in rows], ['scanned', 'SCAN_ERROR', 'DEAD_NODNS'])
        self.assertEqual(rows[1]['CSP'], 'n/a')
        self.assertEqual(rows[2]['CSP'], '')
        self.assertNotIn('\t', (self.root / 'report.csv').read_text())

    def test_13_html_self_contained_and_correct_counts(self):
        self.run_cli(['-o', 'html'])
        text = (self.root / 'report.html').read_text()
        self.assertIn('1 live hosts scanned, 1 scan errors, 1 dead', text)
        self.assertIn('github.com', text.split('Dead / not serving')[1])
        self.assertNotIn('<script src=', text)
        self.assertNotIn('<link', text)
        self.assertIn('manual assessment required', text)

    def test_14_json_full_headers_and_counts(self):
        self.run_cli(['-o', 'json'])
        report = json.loads((self.root / 'report.json').read_text())
        self.assertEqual(report['live_count'], len(report['live']))
        self.assertEqual(report['dead_count'], len(report['dead']))
        self.assertEqual(report['live'][0]['headers'], scanned()['headers'])
        self.assertEqual(report['live'][0]['_meta'], {'dns': 'resolves', 'http_code': '200', 'verdict': 'live'})

    def test_15_workbook_reconciles(self):
        from openpyxl import load_workbook
        self.run_cli(['-o', 'all'])
        wb = load_workbook(self.root / 'report.xlsx')
        self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames, ['Summary', 'Live findings', 'Not assessed', 'Dead hosts'])
        self.assertEqual(wb['Live findings'].max_row, 2)
        self.assertEqual(wb['Not assessed'].max_row, 2)
        self.assertEqual(wb['Dead hosts'].max_row, 2)
        self.assertEqual(wb['Dead hosts']['A2'].value, 'github.com')
        report = json.loads((self.root / 'report.json').read_text())
        self.assertEqual(wb['Summary']['B5'].value, report['live_count'] + report['dead_count'])
        self.assertEqual(wb['Summary']['B6'].value + wb['Summary']['B7'].value, report['live_count'])
        self.assertEqual(wb['Summary']['B8'].value, report['dead_count'])
        self.assertEqual(wb['Live findings']['L2'].value, 2)
        self.assertEqual(wb['Live findings'].freeze_panes, 'B2')
        self.assertEqual(wb['Live findings'].auto_filter.ref, 'A1:L2')

    def test_16_optional_xlsx_skip(self):
        _, err, _ = self.run_cli(['-o', 'all'], missing_xlsx=True)
        self.assertIn('pip install openpyxl', err)
        self.assertIn('4 written, 1 skipped', err)
        self.assertFalse((self.root / 'report.xlsx').exists())
        self.assertEqual(len(list(self.root.iterdir())), 4)

    def test_17_progress_stderr_only(self):
        out, err, _ = self.run_cli()
        self.assertEqual(out, '')
        for marker in ('[1/3]', '[2/3]', '[3/3]', '[  3/3]', '[  2/2]', 'gate done', 'scan done', 'writing done'):
            self.assertIn(marker, err)

    def test_18_quiet_same_reports(self):
        from openpyxl import load_workbook
        self.run_cli(['-o', 'all'])
        original = {p.name: p.read_bytes() for p in self.root.iterdir() if p.suffix != '.xlsx'}
        wb = load_workbook(self.root / 'report.xlsx')
        values = [[list(row) for row in sheet.values] for sheet in wb]
        wb.close()
        out, err, _ = self.run_cli(['-o', 'all'], quiet=True)
        self.assertEqual((out, err), ('', ''))
        self.assertEqual(original, {p.name: p.read_bytes() for p in self.root.iterdir() if p.suffix != '.xlsx'})
        wb = load_workbook(self.root / 'report.xlsx')
        self.assertEqual(values, [[list(row) for row in sheet.values] for sheet in wb])
        wb.close()

    def test_malformed_json_shapes_and_spawn_failure(self):
        for value in (None, [], 1, 'text', {'headers': None}, {'headers': {'server': None}}, {'headers': {}, 'https': None}):
            with self.subTest(value=value), patch.object(hp.subprocess, 'run', return_value=Mock(returncode=0, stdout=json.dumps(value), stderr='')):
                self.assertIn('error', hp.scan_one('example.com'))
        with patch.object(hp.subprocess, 'run', side_effect=FileNotFoundError('missing')):
            self.assertIn('cannot run secheaders', hp.scan_one('example.com')['error'])

    def test_metadata_follows_submitted_url(self):
        targets = ['https://example.com/a', 'https://example.com/b']
        with patch.object(hp, 'check_live', side_effect=lambda t, timeout: ('live', 'resolves', '200' if t.endswith('/a') else '201')), \
             patch.object(hp, 'scan_one', side_effect=lambda t, timeout: scanned('https://github.com/')), \
             contextlib.redirect_stderr(io.StringIO()):
            hp.main(targets + ['-o', 'json', '--outdir', str(self.root), '--name', 'meta'])
        records = json.loads((self.root / 'meta.json').read_text())['live']
        self.assertEqual({r['_input_target']: r['_meta']['http_code'] for r in records}, dict(zip(targets, ['200', '201'])))

    def test_all_4xx_5xx_with_parseable_headers_are_not_assessed(self):
        for status in range(400, 600):
            code = str(status)
            with patch.object(hp, 'check_live', return_value=('live', 'resolves', code)), \
                 patch.object(hp, 'scan_one', return_value=scanned()), contextlib.redirect_stderr(io.StringIO()):
                hp.main(['example.com', '-o', 'json', '--name', 'challenge', '--outdir', str(self.root)])
            record = json.loads((self.root / 'challenge.json').read_text())['live'][0]
            self.assertTrue(record['_scanerror'], code)
            self.assertEqual(record['_meta']['http_code'], code)
            self.assertIn('HTTP ' + code, record['error'])
            self.assertEqual(record['_problems'], 0)
            self.assertEqual(record['_scanner_output']['headers'], scanned()['headers'])

    def test_4xx_5xx_routing_across_all_formats(self):
        from openpyxl import load_workbook
        targets = ['https://example.com/' + str(code) for code in (200, 301, 399, 401, 404, 502)]
        with patch.object(hp, 'check_live', side_effect=lambda t, timeout: ('live', 'resolves', t.rsplit('/', 1)[1])), \
             patch.object(hp, 'scan_one', side_effect=lambda t, timeout: scanned(t)):
            hp.main(targets + ['-o', 'all', '-q', '--outdir', str(self.root), '--name', 'routing'])
        report = json.loads((self.root / 'routing.json').read_text())
        self.assertEqual(report['live_count'], 6)
        self.assertEqual(report['dead_count'], 0)
        self.assertEqual({r['_meta']['http_code'] for r in report['live'] if not r['_scanerror']}, {'200', '301', '399'})
        with (self.root / 'routing.csv').open() as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual(sum(r['scan_status'] == 'SCAN_ERROR' for r in rows), 3)
        for r in rows:
            if int(r['http_code']) >= 400:
                self.assertEqual(r['CSP'], 'n/a')
                self.assertEqual(r['problems'], '')
        self.assertIn('3 live hosts scanned, 3 scan errors, 0 dead', (self.root / 'routing.html').read_text())
        self.assertEqual((self.root / 'routing.txt').read_text().count('SCAN ERROR:'), 3)
        wb = load_workbook(self.root / 'routing.xlsx')
        self.assertEqual(wb['Live findings'].max_row, 4)
        self.assertEqual(wb['Not assessed'].max_row, 4)
        self.assertEqual(wb['Summary']['B6'].value, 3)
        self.assertEqual(wb['Summary']['B7'].value, 3)
        self.assertEqual({r[1] for r in list(wb['Not assessed'].values)[1:]}, {'401', '404', '502'})
        wb.close()

    def test_nondefault_port_is_not_wrong_service_findings(self):
        with patch.object(hp.subprocess, 'run') as scanner:
            self.assertIn('non-default', hp.scan_one('https://example.com:8443/')['error'])
            scanner.assert_not_called()

    def test_cli_validation(self):
        for args in (['example.com', '--workers', '0'], ['example.com', '--timeout', '-1'],
                     ['https://example.com:bad'], ['example.com', '--name', '../report']):
            with self.subTest(args=args), contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                hp.main(args)

    def test_both_worker_pools_overlap(self):
        gate_barrier = threading.Barrier(2, timeout=5)
        scan_barrier = threading.Barrier(2, timeout=5)
        def gate(target, timeout):
            gate_barrier.wait()
            return ('live', 'resolves', '200')
        def scan(target, timeout):
            scan_barrier.wait()
            return scanned(target)
        with patch.object(hp, 'check_live', side_effect=gate), patch.object(hp, 'scan_one', side_effect=scan):
            hp.main(['example.com', 'github.com', '--workers', '2', '-q', '--outdir', str(self.root)])

    def test_subprocess_arguments_and_deadline(self):
        with patch.object(hp.shutil, 'which', return_value=None), \
             patch.object(hp.subprocess, 'run', return_value=Mock(returncode=0, stdout=json.dumps(scanned()), stderr='')) as run:
            hp.scan_one('example.com', timeout=7)
        self.assertEqual(run.call_args.args[0], [sys.executable, '-m', 'secheaders', '--json', 'example.com'])
        self.assertEqual(run.call_args.kwargs['timeout'], 27)
        self.assertNotIn('shell', run.call_args.kwargs)

    def test_xlsx_error_is_literal_text(self):
        from openpyxl import load_workbook
        record = {'target': 'example.com', 'error': '=1+1', '_meta': {'http_code': '200'}}
        hp.render_xlsx(hp.build([record]), [], self.root / 'literal.xlsx', STAMP)
        wb = load_workbook(self.root / 'literal.xlsx')
        self.assertEqual(wb['Not assessed']['C2'].value, '=1+1')
        self.assertEqual(wb['Not assessed']['C2'].data_type, 's')
        wb.close()


if __name__ == '__main__':
    unittest.main()
